from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path(__file__).parent
doc=Document()
s=doc.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(1.8); s.bottom_margin=Cm(1.8); s.left_margin=Cm(2); s.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(30)
doc.styles['Heading 1'].font.size=Pt(21)
doc.styles['Heading 2'].font.size=Pt(13)
doc.core_properties.title='Web de boda de Laura y Javier'
doc.core_properties.subject='Brief de diseño y especificaciones de desarrollo'
def p(t,style=None): return doc.add_paragraph(t,style)
def h(t): doc.add_heading(t,2)
def page(t): doc.add_page_break(); doc.add_heading(t,1)
def bullets(items):
 for x in items: p(x,'List Bullet')
def pic(path,width,alt):
 r=doc.add_paragraph().add_run(); sh=r.add_picture(path,width=Cm(width)); sh._inline.docPr.set('descr',alt)
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Cm(w)
 for c,x in zip(t.rows[0].cells,headers): c.text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row): c.text=x
 for i,row in enumerate(t.rows):
  for c,w in zip(row.cells,widths):
   c.width=Cm(w); pr=c._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); b.append(e)
   pr.append(b)
   if i==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E9E9E9'); pr.append(sh)
   for pa in c.paragraphs:
    pa.paragraph_format.space_after=Pt(5); pa.paragraph_format.space_before=Pt(5)
    for run in pa.runs: run.font.size=Pt(10); run.bold=i==0
  trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

imgs=[r'C:/Users/JAVIER~1.FER/AppData/Local/Temp/codex-clipboard-'+x+'.png' for x in ['83893b33-4a15-4477-bedb-55e513e92749','53f0af39-3627-4b56-bcd2-6d99fb6484b5','836313ba-06ba-45bb-81c4-e95838616b12']]
p('Web de boda de\nLaura y Javier','Title')
p('Diseño y especificaciones para su desarrollo','Subtitle')
pic(imgs[0],15,'Laura y Javier juntos en la playa al atardecer')
p('Una invitación digital cálida y personal, pensada para el móvil, que reúna la información de la boda, la confirmación de asistencia, una playlist compartida y un pequeño juego protagonizado por vuestro perrito.')
p('Este documento sirve de encargo para diseñar y desarrollar la web con Vue 3, TypeScript, Vite, Tailwind CSS y Supabase. Define la experiencia, el contenido, la arquitectura y las condiciones que debe cumplir la entrega.')
h('La idea central')
p('Vuestra historia se recorre siguiendo huellas. Las fotografías y la estética de papel y acuarela aportan emoción; el perrito en pixel art aparece en momentos concretos, acompaña los recuerdos y conduce al juego.')
p('Los textos narrativos y las funciones adicionales se presentan como propuestas. Fecha, lugares, horarios y nombre del perrito quedarán como campos editables hasta que los confirméis.')

page('01 Dirección visual y fotografías')
p('Inspiración de composición: la página de Celia y Mario de Invitify, con portada fotográfica amplia, nombres y fecha centrados, cuenta atrás y recorrido vertical por la información del evento. La adaptación tendrá identidad propia, con vuestras fotos y textos.')
table(['Elemento','Propuesta'],[
('Paleta','Crema #F5F0E7 de fondo; terracota #A85F3D en botones; oliva #72785C en detalles; marrón #49372B para texto.'),
('Tipografía','Serif elegante para títulos y sans serif legible para textos. Caligrafía solo en pequeños acentos, con fuentes de uso permitido.'),
('Composición','Foto de portada con capa oscura suave, espacios amplios, secciones claras y ornamentos botánicos discretos.'),
('Pixel art','Reservado al perrito, sus huellas y el juego. No convertir toda la invitación en una interfaz de videojuego.')],[3.1,13.9])
h('Uso de las imágenes aportadas')
p('Foto de pareja: portada y cierre. En móvil, ajustar el recorte para conservar ambos rostros. Preparar versiones optimizadas y solicitar el original antes de ampliar a pantalla completa.')
pp=doc.add_paragraph(); pp.add_run().add_picture(imgs[1],height=Cm(5.2)); pp.add_run('    '); pp.add_run().add_picture(imgs[2],height=Cm(5.2))
p('Perrito: referencia de pelaje marrón rizado, orejas caídas y silueta para el personaje. Casa y llaves: imagen para el recuerdo del hogar, si confirma la pareja ese hito. Conservar la frase visible como texto alternativo accesible cuando se utilice.')
h('Detalles de interacción')
p('Botones de al menos 44 px, foco visible y contraste comprobado. Animaciones lentas y breves, desactivables con la preferencia de movimiento reducido. Sin música automática ni sonidos al abrir la página.')

page('02 Estructura de la invitación')
table(['Orden','Sección y contenido'],[
('01','Portada con Laura y Javier, fecha pendiente, cuenta atrás y botón Confirmar asistencia.'),
('02','Bienvenida breve y foto. Propuesta: Queremos celebrar contigo el siguiente paso de nuestra historia.'),
('03','Nuestra historia con fotografías, fechas y huellas que conectan cada recuerdo.'),
('04','Ceremonia y celebración con ubicación, hora y enlace al mapa. Añadir al calendario cuando los datos estén confirmados.'),
('05','Programa del día con huellas entre ceremonia, cóctel, comida o cena, baile y fiesta. Todos los horarios serán editables.'),
('06','Información práctica sobre transporte, alojamiento, vestimenta y preguntas frecuentes. Ocultar apartados vacíos.'),
('07','Playlist de la boda con formulario para proponer canciones y botón para abrir Spotify.'),
('08','Minijuego del perrito con instrucciones, opción de jugar y opción de continuar sin jugar.'),
('09','Confirmación de asistencia y mensaje de recepción. Cierre con contacto autorizado de la pareja.')],[1.4,15.6])
h('Dos recorridos con huellas')
p('Nuestra historia contará recuerdos de pareja; el programa del día explicará la boda. Ambos usarán huellas en lugar de flechas, pero tendrán títulos y estilos diferenciados para evitar confusiones.')
p('En móvil, los hitos irán en una columna con huellas a un lado. En escritorio podrán alternarse a ambos lados del camino. El perrito avanzará al aparecer cada hito; las huellas decorativas no serán el único indicador de orden.')
h('Contenido propuesto para la historia')
p('Cómo empezó todo → Aventuras juntos → Nuestro hogar → Nuestra familia con cuatro patas → Nos casamos. Sustituir estos ejemplos por recuerdos reales, sin inventar fechas ni el orden de vuestra historia. En la web, los conectores serán huellas.')

page('03 Playlist de Spotify y asistencia')
h('Una canción para la fiesta')
p('El invitado pega un enlace de canción de Spotify o escribe título y artista. Ve la selección antes de enviarla y recibe un estado claro: propuesta recibida, pendiente de añadir o añadida a Spotify. Un reproductor incrustado sirve para escuchar; no sustituye al formulario de envío.')
p('Propuesta inicial: guardar las sugerencias en Supabase y permitir a Laura y Javier aprobarlas desde el panel. Tras la aprobación, el servidor añade la canción a una playlist de la cuenta organizadora. Los invitados no necesitan conectar su propia cuenta de Spotify.')
bullets(['La cuenta propietaria autoriza la integración mediante OAuth. Solicitar el permiso de modificación que corresponda a la visibilidad de la playlist.',
'Guardar los tokens únicamente en el servidor y renovar el acceso cuando proceda. Nunca incluir secretos en variables VITE_ ni en el navegador.',
'Usar una cola de sincronización, evitar duplicados por identificador de canción y reconciliar el estado tras un fallo ambiguo. Tratar expiración de credenciales, límites de peticiones y falta de permisos.',
'Si la API no está disponible, mantener la propuesta guardada y ofrecer gestión manual. No mostrar Añadida a Spotify hasta confirmar el resultado real.'])
p('Antes de implementar la conexión, comprobar las condiciones actuales de la aplicación Spotify y su modo de cuota. La búsqueda integrada se habilitará solo si está disponible; pegar enlaces o indicar título y artista será siempre una alternativa.')
h('Confirmación de asistencia')
p('Acceso mediante invitación individual o familiar con código no predecible. Campos: asistencia sí o no, asistentes autorizados, necesidades de menú opcionales y transporte si existe. Mostrar los campos de cada asistente únicamente cuando correspondan.')
p('Validar en el servidor la invitación, el número máximo de acompañantes y la fecha límite. Permitir corregir la respuesta desde el mismo acceso, sin generar registros duplicados. Mostrar confirmación solo después de guardar; ante error, conservar el formulario y permitir reintentar.')
p('El panel mostrará respuestas y totales y permitirá exportar un CSV. Las necesidades alimentarias y los datos de contacto no aparecerán en la parte pública. Definir un plazo de conservación y evitar recoger información que no se vaya a utilizar.')

page('04 El perrito y el minijuego')
h('Personaje propio en pixel art')
p('Crear un sprite inspirado en el perrito de la imagen, con pelaje marrón, volumen rizado y orejas caídas. El nombre se mostrará como [NOMBRE DEL PERRITO] hasta confirmarlo. Propuesta de accesorio: pañuelo terracota, sujeto a vuestra elección.')
p('Entregables gráficos previstos: personaje de 64 × 64 px con fondo transparente; animaciones de reposo, caminar y celebrar; huellas de 16 × 16 px; corazón y anillos. Escalar por múltiplos enteros y mantener los bordes nítidos. La fotografía es una referencia, no un sprite terminado.')
h('Juego propuesto Los anillos de Laura y Javier')
p('El perrito debe encontrar dos anillos y llevarlos hasta la pareja. Una partida dura aproximadamente un minuto y transcurre en un pequeño jardín de pixel art, con detalles del hogar de las fotos como decoración.')
bullets(['Inicio: instrucciones breves y botón Jugar. Se carga el juego solo cuando se abre.',
'Control: teclas de dirección o WASD en ordenador; cruceta táctil en móvil. Botones para pausar, reiniciar y salir.',
'Objetivo: recoger dos anillos situados en el escenario y alcanzar a Laura y Javier. Las flores y bancos forman un recorrido sencillo, sin penalizaciones agresivas.',
'Final: el perrito celebra y aparece el mensaje propuesto Los anillos están a salvo. ¡Nos vemos en la boda! Se puede volver a jugar o regresar a la invitación.',
'Accesibilidad: modo tranquilo sin límite de tiempo y alternativa de encontrar los dos anillos mediante botones accesibles. La información de la boda nunca depende de ganar.'])
h('Alcance del juego')
p('Implementación ligera con Canvas 2D, controlada desde un componente Vue. Estados: preparado, jugando, pausado y completado. Pausar al cambiar de pestaña y liberar el bucle de animación al cerrar. Mantener el audio apagado por defecto.')
p('La primera versión no necesita clasificación ni datos personales. Guardar el progreso solo en el dispositivo si se desea. Un ranking sería una ampliación que requeriría validación del resultado en servidor y moderación de alias.')
h('Apariciones del perrito fuera del juego')
p('Saludo junto a la bienvenida, paseo junto a los hitos de la historia y pequeña celebración tras confirmar asistencia. Evitar que siga el cursor o tape botones, formularios o contenido.')

page('05 Arquitectura y datos')
p('Frontend con Vue 3 y Composition API, componentes en TypeScript, Vite y Tailwind CSS. Supabase aporta PostgreSQL, Auth para el panel, Storage para imágenes y funciones de servidor para operaciones privadas e integración con Spotify.')
h('Organización de la aplicación')
p('Ruta principal / para la invitación; /admin para gestión privada. Componentes: WeddingHero, PawTimeline, EventLocations, DaySchedule, GuestInfo, SpotifySuggestions, DogMiniGame y RsvpForm. Separar servicios de datos, validaciones y lógica del juego de los componentes visuales.')
table(['Entidad','Contenido y acceso previsto'],[
('wedding_settings','Nombres, fecha con zona horaria, ubicaciones y textos. Lectura de campos publicados; escritura de administradores.'),
('story_milestones','Título, fecha opcional, texto, imagen, orden y estado publicado. Solo hitos publicados visibles.'),
('invitations y guests','Invitación, hash del código, cupo y asistentes permitidos. Acceso privado por servidor; nunca listados públicos.'),
('rsvps','Respuesta por invitación y asistente, menú opcional, transporte y fecha de actualización. Acceso por invitación validada y administradores.'),
('song_suggestions','Enlace o identificador, título, artista y estado de aprobación. Evitar duplicados; escritura validada y moderación privada.'),
('spotify_sync_jobs','Canción, estado, intentos y resultado de sincronización. Solo servidor; credenciales fuera de tablas públicas.')],[4.1,12.9])
h('Autorización y protección de datos')
p('Activar RLS en las tablas expuestas y conceder solo las operaciones necesarias. La clave pública del frontend no sustituye las políticas de acceso. Las credenciales privilegiadas permanecen exclusivamente en el servidor. [2]')
p('Validar el código de invitación en una función de servidor y vincularlo a una sesión limitada a esa invitación. Aplicar caducidad, límites de intentos y comprobación en cada escritura. No permitir consultar o modificar otra familia cambiando un identificador.')
p('Administradores autenticados con rol asignado por servidor; ocultar la ruta no basta. Almacenar fotografías privadas en un bucket privado con acceso autorizado. Validar campos y enlaces, limitar peticiones y evitar registrar códigos, tokens o información alimentaria en logs.')

page('06 Entrega y criterios de aceptación')
h('Orden de desarrollo')
p('1. Confirmar contenidos y aprobar portada, paleta y recorrido móvil. 2. Construir la invitación adaptable y las huellas. 3. Conectar asistencia y panel. 4. Integrar propuestas musicales y sincronización. 5. Incorporar el perrito y el juego. 6. Probar, completar los datos y publicar en el alojamiento elegido.')
h('Comprobaciones de entrega')
bullets(['La portada conserva los rostros y el texto se lee en móvil y escritorio. No hay desplazamiento horizontal desde 360 px de ancho.',
'La cuenta atrás usa la fecha y zona horaria configuradas y cambia de mensaje al llegar el día. Los mapas y el calendario contienen los datos aprobados.',
'Las huellas conectan los hitos y la lectura sigue siendo clara con teclado, lector de pantalla y movimiento reducido.',
'La asistencia se puede guardar y corregir, respeta el cupo y rechaza una invitación no válida. Una familia no puede leer ni modificar otra.',
'Las canciones no se duplican y el estado refleja la sincronización real. Probar permisos revocados, token caducado, límite de API y pérdida de red.',
'El juego se puede completar en móvil y ordenador, pausar y cerrar. La alternativa accesible permite llegar al mismo mensaje final.',
'El panel requiere autorización real. No hay secretos en el paquete del navegador ni respuestas privadas accesibles sin permiso.',
'La compilación de producción y la comprobación TypeScript terminan correctamente. Entregar código, instrucciones, variables de ejemplo, migraciones y recursos gráficos.'])
h('Datos por completar antes de publicar')
p('Fecha y hora con zona horaria; lugares y enlaces de mapa; horarios y transporte; plazo de confirmación; invitados y cupos; nombre del perrito; recuerdos y fechas reales; cuenta y playlist de Spotify; fotos originales; contacto autorizado; dominio y alojamiento.')
h('Referencias para el desarrollo')
p('Consultadas el 16 de octubre de 2026. Revisar la documentación vigente al implementar las integraciones.')
for t in [
'[1] Referencia visual de Invitify: https://invitify.es/bodas/celia-y-mario/',
'[2] Supabase y políticas RLS: https://supabase.com/docs/guides/database/postgres/row-level-security',
'[3] Spotify y alta de canciones: https://developer.spotify.com/documentation/web-api/reference/add-items-to-playlist',
'[4] Spotify y autorización: https://developer.spotify.com/documentation/web-api/concepts/authorization',
'[5] Spotify y modo de cuota: https://developer.spotify.com/documentation/web-api/concepts/quota-modes']:
 pa=p(t)
 for r in pa.runs: r.font.size=Pt(9)
doc.save(ROOT/'Web_boda_Laura_y_Javier.docx')
print(ROOT/'Web_boda_Laura_y_Javier.docx')
